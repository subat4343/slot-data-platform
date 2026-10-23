"""Phase 6: 出力・監査・バックアップ復元・CLI。"""

from __future__ import annotations

import csv
import os
from datetime import date
from pathlib import Path

import pytest
from sqlalchemy import create_engine, func, select, text

from slot_data.exports.export import ExportFilter, bonus_rows, daily_rows, slump_rows, write_rows
from slot_data.ingestion.pipeline import IngestionPipeline
from slot_data.models import DailyMachineStat, MachineModel, RawPayload, SlumpPoint
from slot_data.quality.audit import run_audit
from tests.conftest import ROOT, TEST_DB_URL
from tests.helpers import manual_doc, write_doc

pytestmark = pytest.mark.integration


def _load_sample(ctx, tmp_path):
    with ctx.session() as s:
        s.add(MachineModel(canonical_name="機種A"))
    doc = manual_doc(
        [
            {
                "machine_number": "301",
                "model_name": "機種A",
                "fields": {"total_games": 5000, "net_medals": -300, "bb_count": 12, "rb_count": None},
                "bonus_events": {
                    "status": "complete",
                    "events": [{"ordinal": 1, "hit_games": 100, "raw_label": "BIG"}],
                },
                "slump_series": [
                    {
                        "x_axis_type": "cumulative_games",
                        "x_unit": "games",
                        "y_axis_type": "net_medals",
                        "y_unit": "medals",
                        "points": [{"x_numeric": i, "y_numeric": i * 2} for i in range(50)],
                    }
                ],
            },
            {"machine_number": "302", "model_name": "機種A", "fields": {"total_games": 0}},
        ]
    )
    IngestionPipeline(ctx).import_file("pscube", "H001", write_doc(tmp_path, "s.json", doc))


@pytest.mark.parametrize("fmt", ["csv", "xlsx", "parquet"])
def test_export_formats_keep_nulls_and_sources(ctx, hall, tmp_path, fmt):
    _load_sample(ctx, tmp_path)
    out = tmp_path / f"daily.{fmt}"
    with ctx.session() as s:
        count = write_rows(
            daily_rows(s, ExportFilter(hall_codes=["H001"], model_name="機種"), with_layout=True), out, fmt
        )
    assert count == 2
    if fmt == "csv":
        rows = list(csv.DictReader(out.open(encoding="utf-8-sig")))
        r301 = next(r for r in rows if r["machine_number"] == "301")
        assert r301["rb_count"] == "" and r301["rb_count__status"] == "not_published"
        assert r301["source_code"] == "pscube" and r301["weekday"] == "火"
        r302 = next(r for r in rows if r["machine_number"] == "302")
        assert r302["total_games"] == "0" and r302["total_games__status"] == "observed"
    elif fmt == "parquet":
        import pyarrow.parquet as pq

        table = pq.read_table(out)
        assert table.num_rows == 2 and "net_medals__status" in table.column_names
    else:
        from openpyxl import load_workbook

        ws = load_workbook(out).active
        assert ws.max_row == 3


def test_export_filters_and_history(ctx, hall, tmp_path):
    _load_sample(ctx, tmp_path)
    with ctx.session() as s:
        assert len(list(daily_rows(s, ExportFilter(machine_numbers=["302"])))) == 1
        assert len(list(daily_rows(s, ExportFilter(date_from=date(2026, 10, 21))))) == 0
        assert len(list(daily_rows(s, ExportFilter(model_name="存在しない")))) == 0
        assert len(list(bonus_rows(s, ExportFilter()))) == 1
        n = write_rows(
            slump_rows(s, ExportFilter(hall_codes=["H001"])), tmp_path / "slump.parquet", "parquet"
        )
        assert n == 50 == s.scalar(select(func.count()).select_from(SlumpPoint))


def test_audit_detects_missing_and_orphan_raw(ctx, hall, tmp_path):
    _load_sample(ctx, tmp_path)
    report = run_audit(ctx)
    assert report.ok and report.raw_checked == 1
    with ctx.session() as s:
        uri = s.scalar(select(RawPayload.storage_uri))
    path = ctx.raw_store.path_for(uri)
    os.chmod(path, 0o666)
    path.unlink()
    ctx.raw_store.put(b"orphan bytes")
    report = run_audit(ctx)
    assert not report.ok
    assert report.raw_missing == [uri] and len(report.orphan_files) == 1


def _pg_bin() -> Path | None:
    env = os.environ.get("SLOT_DATA_PG_BIN_DIR")
    candidate = Path(env) if env else ROOT / ".devdb" / "pgsql" / "bin"
    return candidate if candidate.exists() else None


def test_backup_and_restore_roundtrip(ctx, hall, tmp_path):
    from slot_data.storage.backup import BackupError, backup, restore
    from slot_data.storage.raw_store import LocalRawStore

    bin_dir = _pg_bin()
    _load_sample(ctx, tmp_path)
    try:
        result = backup(TEST_DB_URL, ctx.raw_store, tmp_path / "bk", bin_dir)
    except BackupError as exc:
        pytest.skip(f"pg_dump unavailable: {exc}")
    assert result.raw_files == 1

    restore_db = "slot_data_restore_test"
    admin = create_engine(TEST_DB_URL, isolation_level="AUTOCOMMIT")
    with admin.connect() as conn:
        conn.execute(text(f"DROP DATABASE IF EXISTS {restore_db}"))
        conn.execute(text(f"CREATE DATABASE {restore_db}"))
    restore_url = TEST_DB_URL.rsplit("/", 1)[0] + f"/{restore_db}"
    restored_store = LocalRawStore(tmp_path / "restored-raw")
    info = restore(tmp_path / "bk", restore_url, restored_store, bin_dir)
    assert info["raw_verified"] == 1

    engine = create_engine(restore_url)
    with engine.connect() as conn:
        assert conn.execute(select(func.count()).select_from(DailyMachineStat)).scalar_one() == 2
        uri = conn.execute(select(RawPayload.storage_uri)).scalar_one()
        # 追記専用トリガーも復元されている
        with pytest.raises(Exception, match="append-only"):
            conn.execute(text("DELETE FROM raw_payloads"))
    engine.dispose()
    assert restored_store.get(uri) == ctx.raw_store.get(uri)
    with admin.connect() as conn:
        conn.execute(text(f"DROP DATABASE IF EXISTS {restore_db}"))
    admin.dispose()


def test_cli_end_to_end(engine, session_factory, tmp_path, monkeypatch):
    from typer.testing import CliRunner

    from slot_data import config as config_mod
    from slot_data.cli.main import app

    monkeypatch.setenv("SLOT_DATA_DATABASE_URL", TEST_DB_URL)
    monkeypatch.setenv("SLOT_DATA_RAW_STORE_ROOT", str(tmp_path / "raw"))
    monkeypatch.setenv("SLOT_DATA_CONFIG_PATH", str(ROOT / "config" / "app.example.yaml"))
    config_mod.get_settings.cache_clear()
    config_mod.get_app_config.cache_clear()
    runner = CliRunner()

    def run(*args, code=0):
        result = runner.invoke(app, ["--text-log", *args])
        assert result.exit_code == code, result.output + repr(result.exception)
        return result.output

    run("hall", "add", "H001", "テスト店", "--cutoff", "05:00", "--closing", "22:45")
    obs = write_doc(
        tmp_path,
        "o.json",
        manual_doc([{"machine_number": "301", "model_name": "機種X", "fields": {"total_games": 10}}]),
    )
    run("import-observations", str(obs), "--source", "pscube", "--hall", "H001")
    pending = run("mapping", "list", "--pending")
    assert "機種X" in pending
    import json

    mapping_id = json.loads(pending)[0]["id"]
    run("mapping", "resolve", mapping_id, "--create-model", "機種X")
    run("reconcile", "--hall", "H001", "--from", "2026-10-20", "--to", "2026-10-20")
    run(
        "correction",
        "add",
        "--hall",
        "H001",
        "--date",
        "2026-10-20",
        "--machine",
        "301",
        "--field",
        "bb_count",
        "--value",
        "3",
        "--reason",
        "目視確認",
        "--by",
        "tester",
    )
    layout = tmp_path / "layout.csv"
    layout.write_text(
        "hall_code,layout_valid_from,island_code,side,sequence_no,machine_number\n"
        "H001,2026-10-01,A,L,1,301\nH001,2026-10-01,A,L,2,302\n",
        encoding="utf-8",
    )
    run("validate-layout", str(layout))
    staged = json.loads(run("import-layout", str(layout), "--by", "tester").split("確認後に公開")[0])
    run("publish-layout", staged["import_id"], "--by", "tester", code=1)  # --yes 無しは拒否
    run("publish-layout", staged["import_id"], "--by", "tester", "--yes")
    shown = json.loads(run("layout-show", "H001", "--date", "2026-10-20"))
    assert shown["301"]["corner_rank"] == 1 and shown["301"]["next_to"] == ["302"]
    out = tmp_path / "daily.csv"
    run("export", "daily", "--out", str(out), "--hall", "H001", "--with-layout")
    rows = list(csv.DictReader(out.open(encoding="utf-8-sig")))
    assert rows[0]["model_name"] == "機種X" and rows[0]["bb_count"] == "3" and rows[0]["corner_rank"] == "1"
    run("audit")
    run("issues")
    # 自動収集は利用条件により拒否される（exit 2）
    monkeypatch.setenv("SLOT_DATA_CONFIG_PATH", str(ROOT / "config" / "app.example.yaml"))
    out = run("collect", "--source", "pscube", "--hall", "H001", "--date", "2026-10-20", code=2)
    config_mod.get_settings.cache_clear()
    config_mod.get_app_config.cache_clear()
