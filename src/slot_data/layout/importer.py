"""配置 CSV/Excel の取り込み: validate -> preview -> confirm -> publish。

- ``validate_rows``: 形式・重複・参照・欠番を検証（DB は参照のみ）
- ``stage_import``: 検証結果を layout_imports に記録し、正常なら draft の配置版を作って差分プレビューを返す
- ``publish_import``: 人の確認後に draft を原子的に公開（前の公開版の valid_to を確定）
不正なファイルは draft を作らず、公開版には一切反映しない。
"""

from __future__ import annotations

import csv
import hashlib
import io
import uuid
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any

from sqlalchemy import func, select, text
from sqlalchemy.orm import Session

from slot_data.domain.enums import LayoutStatus, PositionKind, RelationType
from slot_data.domain.timeutil import utcnow
from slot_data.layout.geometry import Slot, derive_relations, sequence_gaps
from slot_data.models import (
    Hall,
    HallIsland,
    HallLayoutVersion,
    LayoutImport,
    MachineAdjacency,
    MachinePosition,
)
from slot_data.schemas.parsed import normalize_machine_number

REQUIRED = ("hall_code", "layout_valid_from", "island_code", "side", "sequence_no", "machine_number")
OPTIONAL = ("x", "y", "rotation_deg", "position_kind", "notes")


class LayoutError(ValueError):
    pass


@dataclass
class LayoutRow:
    line: int
    hall_code: str
    valid_from: date
    island_code: str
    side: str
    sequence_no: int
    machine_number: str | None
    position_kind: str
    x: Decimal | None = None
    y: Decimal | None = None
    rotation_deg: Decimal | None = None
    notes: str | None = None


@dataclass
class ValidationReport:
    errors: list[dict[str, Any]] = field(default_factory=list)
    warnings: list[dict[str, Any]] = field(default_factory=list)
    rows: list[LayoutRow] = field(default_factory=list)

    @property
    def ok(self) -> bool:
        return not self.errors

    def error(self, code: str, message: str, line: int | None = None) -> None:
        self.errors.append({"code": code, "message": message, "line": line})

    def warning(self, code: str, message: str, line: int | None = None) -> None:
        self.warnings.append({"code": code, "message": message, "line": line})

    def to_json(self) -> dict[str, Any]:
        return {"ok": self.ok, "row_count": len(self.rows), "errors": self.errors, "warnings": self.warnings}


def read_table(path: Path) -> list[dict[str, str]]:
    suffix = path.suffix.lower()
    if suffix == ".csv":
        raw = path.read_bytes()
        try:
            content = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            content = raw.decode("cp932")
        return [dict(r) for r in csv.DictReader(io.StringIO(content))]
    if suffix in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = [str(h).strip() if h is not None else "" for h in next(rows)]
        out = []
        for values in rows:
            if values is None or all(v is None for v in values):
                continue
            out.append(
                {
                    h: ("" if v is None else (v.isoformat()[:10] if hasattr(v, "isoformat") else str(v)))
                    for h, v in zip(header, values, strict=False)
                }
            )
        wb.close()
        return out
    raise LayoutError(f"unsupported layout file: {path.name}（.csv / .xlsx）")


def _decimal(value: str, name: str, line: int, report: ValidationReport) -> Decimal | None:
    if value is None or str(value).strip() == "":
        return None
    try:
        return Decimal(str(value).strip())
    except InvalidOperation:
        report.error("INVALID_NUMBER", f"{name}={value!r} is not numeric", line)
        return None


def validate_rows(session: Session, records: list[dict[str, str]]) -> ValidationReport:
    report = ValidationReport()
    if not records:
        report.error("EMPTY", "no rows")
        return report
    missing = [c for c in REQUIRED if c not in records[0]]
    if missing:
        report.error("MISSING_COLUMNS", f"required columns missing: {missing}")
        return report
    unknown = set(records[0]) - set(REQUIRED) - set(OPTIONAL)
    if unknown:
        report.warning("UNKNOWN_COLUMNS", f"ignored columns: {sorted(unknown)}")

    for index, rec in enumerate(records, start=2):

        def get(key: str, rec: dict[str, str] = rec) -> str:
            return (rec.get(key) or "").strip()

        try:
            valid_from = date.fromisoformat(get("layout_valid_from"))
        except ValueError:
            report.error("INVALID_DATE", f"layout_valid_from={get('layout_valid_from')!r}", index)
            continue
        try:
            seq = int(get("sequence_no"))
            if seq < 1:
                raise ValueError
        except ValueError:
            report.error(
                "INVALID_SEQUENCE", f"sequence_no={get('sequence_no')!r} must be an integer >= 1", index
            )
            continue
        kind = get("position_kind") or PositionKind.MACHINE.value
        if kind not in {k.value for k in PositionKind}:
            report.error("INVALID_POSITION_KIND", f"position_kind={kind!r}", index)
            continue
        number_raw = get("machine_number")
        if kind == PositionKind.MACHINE and not number_raw:
            report.error(
                "MISSING_MACHINE_NUMBER", "machine_number is required for position_kind=machine", index
            )
            continue
        if kind != PositionKind.MACHINE and number_raw:
            report.error(
                "UNEXPECTED_MACHINE_NUMBER", f"position_kind={kind} must not have machine_number", index
            )
            continue
        if not get("island_code") or not get("side") or not get("hall_code"):
            report.error("MISSING_VALUE", "hall_code / island_code / side are required", index)
            continue
        report.rows.append(
            LayoutRow(
                line=index,
                hall_code=get("hall_code"),
                valid_from=valid_from,
                island_code=get("island_code"),
                side=get("side").upper(),
                sequence_no=seq,
                machine_number=normalize_machine_number(number_raw) if number_raw else None,
                position_kind=kind,
                x=_decimal(get("x"), "x", index, report),
                y=_decimal(get("y"), "y", index, report),
                rotation_deg=_decimal(get("rotation_deg"), "rotation_deg", index, report),
                notes=get("notes") or None,
            )
        )
    if not report.rows:
        return report

    halls = {r.hall_code for r in report.rows}
    dates = {r.valid_from for r in report.rows}
    if len(halls) != 1:
        report.error("MULTIPLE_HALLS", f"one file must contain one hall: {sorted(halls)}")
    if len(dates) != 1:
        report.error(
            "MULTIPLE_VALID_FROM", f"one file must contain one layout_valid_from: {sorted(map(str, dates))}"
        )

    seen_numbers: dict[str, int] = {}
    seen_slots: dict[tuple[str, str, int], int] = {}
    for row in report.rows:
        if row.machine_number is not None:
            if row.machine_number in seen_numbers:
                report.error(
                    "DUPLICATE_MACHINE_NUMBER",
                    f"machine {row.machine_number} also on line {seen_numbers[row.machine_number]}",
                    row.line,
                )
            seen_numbers.setdefault(row.machine_number, row.line)
        slot = (row.island_code, row.side, row.sequence_no)
        if slot in seen_slots:
            report.error("DUPLICATE_POSITION", f"position {slot} also on line {seen_slots[slot]}", row.line)
        seen_slots.setdefault(slot, row.line)

    sides: dict[str, set[str]] = {}
    for row in report.rows:
        sides.setdefault(row.island_code, set()).add(row.side)
    for island, island_sides in sides.items():
        if len(island_sides) > 2:
            report.warning(
                "MANY_SIDES", f"island {island} has {len(island_sides)} sides; back_to_back not derived"
            )
    for island, side, a, b in sequence_gaps(_slots(report.rows)):
        report.warning(
            "SEQUENCE_GAP",
            f"island {island} side {side}: sequence {a} -> {b}; adjacency not derived across the gap",
        )

    if len(halls) == 1 and len(dates) == 1:
        hall = session.scalar(select(Hall).where(Hall.code == next(iter(halls))))
        if hall is None:
            report.error("UNKNOWN_HALL", f"hall_code {next(iter(halls))} is not registered")
        else:
            valid_from = next(iter(dates))
            latest = session.scalar(
                select(func.max(HallLayoutVersion.valid_from)).where(
                    HallLayoutVersion.hall_id == hall.id, HallLayoutVersion.status == LayoutStatus.PUBLISHED
                )
            )
            if latest is not None and valid_from <= latest:
                report.error(
                    "PERIOD_OVERLAP",
                    f"layout_valid_from {valid_from} must be after the latest published version ({latest})",
                )
    return report


def _slots(rows: list[LayoutRow]) -> list[Slot]:
    return [
        Slot(r.line, r.island_code, r.side, r.sequence_no, r.position_kind, r.machine_number) for r in rows
    ]


def effective_version(session: Session, hall_id: uuid.UUID, day: date) -> HallLayoutVersion | None:
    """その営業日に有効な公開済み配置版。"""
    return session.scalar(
        select(HallLayoutVersion).where(
            HallLayoutVersion.hall_id == hall_id,
            HallLayoutVersion.status == LayoutStatus.PUBLISHED,
            HallLayoutVersion.valid_from <= day,
            (HallLayoutVersion.valid_to.is_(None)) | (HallLayoutVersion.valid_to > day),
        )
    )


def _position_map(session: Session, version_id: uuid.UUID) -> dict[str, tuple[str, str, int]]:
    rows = session.execute(
        select(
            MachinePosition.machine_number,
            HallIsland.island_code,
            MachinePosition.side,
            MachinePosition.sequence_no,
        )
        .join(HallIsland, HallIsland.id == MachinePosition.island_id)
        .where(MachinePosition.layout_version_id == version_id, MachinePosition.machine_number.is_not(None))
    )
    return {r.machine_number: (r.island_code, r.side, r.sequence_no) for r in rows}


def diff_preview(session: Session, hall_id: uuid.UUID, rows: list[LayoutRow]) -> dict[str, Any]:
    new = {r.machine_number: (r.island_code, r.side, r.sequence_no) for r in rows if r.machine_number}
    current_version = effective_version(session, hall_id, rows[0].valid_from)
    old = _position_map(session, current_version.id) if current_version else {}
    by_slot_old = {v: k for k, v in old.items()}
    added = sorted(set(new) - set(old), key=_num_key)
    removed = sorted(set(old) - set(new), key=_num_key)
    moved = sorted((k for k in set(new) & set(old) if new[k] != old[k]), key=_num_key)
    renumbered = [
        {"slot": list(slot), "from": by_slot_old[slot], "to": number}
        for number, slot in new.items()
        if number in added and slot in by_slot_old and by_slot_old[slot] in removed
    ]
    return {
        "base_version": current_version.version_number if current_version else None,
        "added": added,
        "removed": removed,
        "moved": [{"machine_number": k, "from": list(old[k]), "to": list(new[k])} for k in moved],
        "renumbered_candidates": renumbered,
        "unchanged": len(set(new) & set(old)) - len(moved),
        "has_changes": bool(added or removed or moved),
    }


def _num_key(value: str) -> tuple:
    return (0, int(value)) if value.isdigit() else (1, value)


@dataclass
class StagedImport:
    import_id: uuid.UUID
    report: dict[str, Any]
    preview: dict[str, Any] | None
    draft_version_id: uuid.UUID | None


def stage_import(session: Session, path: Path, created_by: str, file_uri: str | None = None) -> StagedImport:
    data = path.read_bytes()
    records = read_table(path)
    report = validate_rows(session, records)
    hall = None
    if report.rows:
        hall = session.scalar(select(Hall).where(Hall.code == report.rows[0].hall_code))
    if hall is None:
        raise LayoutError(f"cannot stage import: {report.to_json()}")
    preview = diff_preview(session, hall.id, report.rows) if report.ok else None
    if report.ok and preview and not preview["has_changes"] and preview["base_version"] is not None:
        report.error(
            "NO_CHANGES", "layout is identical to the current published version; no new version needed"
        )

    imp = LayoutImport(
        hall_id=hall.id,
        file_uri=file_uri or path.name,
        file_sha256=hashlib.sha256(data).hexdigest(),
        validation_report={**report.to_json(), "preview": preview},
        status="previewed" if report.ok else "invalid",
        created_by=created_by,
    )
    session.add(imp)
    session.flush()
    if not report.ok:
        return StagedImport(imp.id, report.to_json(), preview, None)

    version = _create_draft(session, hall, report.rows, created_by, file_uri or path.name)
    imp.target_layout_version_id = version.id
    session.flush()
    return StagedImport(imp.id, report.to_json(), preview, version.id)


def _create_draft(
    session: Session, hall: Hall, rows: list[LayoutRow], created_by: str, uri: str
) -> HallLayoutVersion:
    next_number = (
        session.scalar(
            select(func.max(HallLayoutVersion.version_number)).where(HallLayoutVersion.hall_id == hall.id)
        )
        or 0
    ) + 1
    version = HallLayoutVersion(
        hall_id=hall.id,
        version_number=next_number,
        valid_from=rows[0].valid_from,
        status=LayoutStatus.DRAFT,
        created_by=created_by,
        source_document_uri=uri,
        layout_metadata={
            "sequence_basis": "sequence_no は同じ島の両側で共通の物理的基準端から数える",
            "row_count": len(rows),
        },
    )
    session.add(version)
    session.flush()
    islands: dict[str, HallIsland] = {}
    for code in sorted({r.island_code for r in rows}):
        islands[code] = HallIsland(layout_version_id=version.id, island_code=code, display_name=code)
        session.add(islands[code])
    session.flush()
    positions: dict[int, MachinePosition] = {}
    for row in rows:
        pos = MachinePosition(
            layout_version_id=version.id,
            island_id=islands[row.island_code].id,
            machine_number=row.machine_number,
            side=row.side,
            sequence_no=row.sequence_no,
            x=row.x,
            y=row.y,
            rotation_deg=row.rotation_deg,
            position_kind=row.position_kind,
            notes=row.notes,
        )
        session.add(pos)
        positions[row.line] = pos
    session.flush()
    for rel in derive_relations(_slots(rows)):
        a, b = sorted([positions[rel.a].id, positions[rel.b].id])
        session.add(
            MachineAdjacency(
                layout_version_id=version.id,
                position_a_id=a,
                position_b_id=b,
                relation_type=rel.relation_type,
            )
        )
    session.flush()
    return version


def set_relation(
    session: Session,
    version_id: uuid.UUID,
    machine_a: str,
    machine_b: str,
    relation_type: str,
    remove: bool = False,
    notes: str | None = None,
) -> None:
    """draft 版の位置関係を手動補正する（公開済み版はトリガーで拒否される）。"""
    version = session.get(HallLayoutVersion, version_id)
    if version is None or version.status != LayoutStatus.DRAFT:
        raise LayoutError("relations can only be edited on a draft layout version")
    if relation_type not in {r.value for r in RelationType}:
        raise LayoutError(f"unknown relation type: {relation_type}")
    ids = []
    for number in (machine_a, machine_b):
        pos = session.scalar(
            select(MachinePosition).where(
                MachinePosition.layout_version_id == version_id,
                MachinePosition.machine_number == normalize_machine_number(number),
            )
        )
        if pos is None:
            raise LayoutError(f"machine {number} not in layout version")
        ids.append(pos.id)
    if ids[0] == ids[1]:
        raise LayoutError("a machine cannot relate to itself")
    a, b = sorted(ids)
    existing = session.scalar(
        select(MachineAdjacency).where(
            MachineAdjacency.layout_version_id == version_id,
            MachineAdjacency.position_a_id == a,
            MachineAdjacency.position_b_id == b,
            MachineAdjacency.relation_type == relation_type,
        )
    )
    if remove:
        if existing is not None:
            session.delete(existing)
    elif existing is None:
        session.add(
            MachineAdjacency(
                layout_version_id=version_id,
                position_a_id=a,
                position_b_id=b,
                relation_type=relation_type,
                is_manual_override=True,
                notes=notes,
            )
        )
    session.flush()


def check_relations(session: Session, version_id: uuid.UUID) -> list[str]:
    """矛盾する位置関係（同じ組に複数種別、隣が別の島など）を検出する。"""
    problems = []
    rows = session.execute(
        select(
            MachineAdjacency.position_a_id,
            MachineAdjacency.position_b_id,
            func.array_agg(MachineAdjacency.relation_type),
        )
        .where(MachineAdjacency.layout_version_id == version_id)
        .group_by(MachineAdjacency.position_a_id, MachineAdjacency.position_b_id)
    ).all()
    positions = {
        p.id: p
        for p in session.scalars(
            select(MachinePosition).where(MachinePosition.layout_version_id == version_id)
        )
    }
    for a, b, types in rows:
        pa, pb = positions[a], positions[b]
        label = f"{pa.machine_number}-{pb.machine_number}"
        if len(types) > 1:
            problems.append(f"{label}: multiple relation types {sorted(types)}")
        if RelationType.NEXT_TO in types and (pa.island_id != pb.island_id or pa.side != pb.side):
            problems.append(f"{label}: next_to across different island/side")
        if RelationType.BACK_TO_BACK in types and (pa.island_id != pb.island_id or pa.side == pb.side):
            problems.append(f"{label}: back_to_back must be same island, opposite side")
        if RelationType.ACROSS_AISLE in types and pa.island_id == pb.island_id:
            problems.append(f"{label}: across_aisle must connect different islands")
    return problems


def publish_import(session: Session, import_id: uuid.UUID, confirmed_by: str) -> HallLayoutVersion:
    imp = session.get(LayoutImport, import_id)
    if imp is None:
        raise LayoutError(f"layout import not found: {import_id}")
    if imp.status != "previewed" or imp.target_layout_version_id is None:
        raise LayoutError(f"import {import_id} is not publishable (status={imp.status})")
    session.execute(
        text("SELECT pg_advisory_xact_lock(hashtextextended(:k, 0))"), {"k": f"layout:{imp.hall_id}"}
    )
    version = session.get(HallLayoutVersion, imp.target_layout_version_id, with_for_update=True)
    problems = check_relations(session, version.id)
    if problems:
        raise LayoutError("conflicting relations: " + "; ".join(problems))
    current = session.scalar(
        select(HallLayoutVersion)
        .where(HallLayoutVersion.hall_id == imp.hall_id, HallLayoutVersion.status == LayoutStatus.PUBLISHED)
        .order_by(HallLayoutVersion.valid_from.desc())
        .limit(1)
        .with_for_update()
    )
    if current is not None:
        if current.valid_from >= version.valid_from:
            raise LayoutError(
                f"cannot publish valid_from {version.valid_from}: latest published version starts {current.valid_from}"
            )
        if current.valid_to is None:
            current.valid_to = version.valid_from
        elif current.valid_to > version.valid_from:
            raise LayoutError("latest published version already ends after the new valid_from")
    session.flush()
    version.status = LayoutStatus.PUBLISHED
    version.published_at = utcnow()
    version.notes = (version.notes or "") + f"confirmed_by={confirmed_by}"
    imp.status = "published"
    session.flush()
    return version


def reject_import(session: Session, import_id: uuid.UUID) -> None:
    imp = session.get(LayoutImport, import_id)
    if imp is None:
        raise LayoutError(f"layout import not found: {import_id}")
    if imp.status != "previewed":
        raise LayoutError(f"import {import_id} cannot be rejected (status={imp.status})")
    imp.status = "rejected"
    if imp.target_layout_version_id:
        version = session.get(HallLayoutVersion, imp.target_layout_version_id)
        version.status = LayoutStatus.REJECTED
    session.flush()
